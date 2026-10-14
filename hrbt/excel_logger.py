"""Excel event log.

Layout (sheet "Event Log", one row per date):
    A: Date
    B: Stream 1 events
    C: Stream 1 seconds blocked
    D: Stream 2 events
    E: Stream 2 seconds blocked

The workbook is opened when logging starts, updated in memory as events
happen, saved to disk regularly, and closed when logging stops. If the file is
open in Excel (Windows locks it), the save fails, the error is written to
errors.log and the save is retried on the next cycle - nothing is lost while
the app is running.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from .error_log import log_error

SHEET_NAME = "Event Log"
HEADERS = [
    "Date",
    "Stream 1 Events",
    "Stream 1 Seconds Blocked",
    "Stream 2 Events",
    "Stream 2 Seconds Blocked",
]
# column index (1-based) for (events, seconds) of each stream
_STREAM_COLS = {1: (2, 3), 2: (4, 5)}


class ExcelLogError(RuntimeError):
    pass


def create_workbook(path: Path) -> None:
    """Create a new, empty log file with the header row."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for col, width in zip("ABCDE", (14, 16, 24, 16, 24)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return datetime.strptime(value.strip()[:10], "%Y-%m-%d").date()
        except ValueError:
            return None
    return None


def _as_number(value) -> float:
    try:
        return float(value) if value is not None else 0
    except (TypeError, ValueError):
        return 0


class ExcelLogger:
    def __init__(self, path: Path):
        self.path = Path(path)
        self._wb = None
        self._ws = None
        self._rows: dict[date, int] = {}
        self._dirty = False

    # -------------------------------------------------------------- lifecycle
    @property
    def is_open(self) -> bool:
        return self._wb is not None

    def open(self) -> None:
        if not self.path.exists():
            raise ExcelLogError(f"{self.path.name} does not exist.")
        try:
            wb = load_workbook(self.path)
        except PermissionError as exc:
            raise ExcelLogError(
                f"{self.path.name} is locked - close it in Excel and try again.") from exc
        except Exception as exc:  # corrupt / not an xlsx
            raise ExcelLogError(f"Could not open {self.path.name}: {exc}") from exc

        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]
        if ws.cell(1, 1).value is None:  # empty sheet - restore the header row
            for i, h in enumerate(HEADERS, start=1):
                ws.cell(1, i, h)

        rows: dict[date, int] = {}
        for r in range(2, ws.max_row + 1):
            d = _as_date(ws.cell(r, 1).value)
            if d is not None:
                rows[d] = r
        self._wb, self._ws, self._rows = wb, ws, rows
        self._dirty = False

    def close(self) -> bool:
        """Save and release the workbook.

        Returns False (and keeps the workbook in memory) if the save failed,
        so the caller can retry - or call discard() to give up.
        """
        if not self.is_open:
            return True
        if not self.save(force=True):
            return False
        self.discard()
        return True

    def discard(self) -> None:
        """Release the workbook without saving."""
        if self._wb is not None:
            try:
                self._wb.close()
            finally:
                self._wb = self._ws = None
                self._rows = {}

    def save(self, force: bool = False) -> bool:
        if not self.is_open or (not self._dirty and not force):
            return True
        tmp = self.path.with_name(self.path.stem + ".saving" + self.path.suffix)
        try:
            self._wb.save(tmp)
            tmp.replace(self.path)
            self._dirty = False
            return True
        except PermissionError as exc:
            log_error(f"Could not save {self.path.name} (is it open in Excel?). "
                      "Will retry automatically.", exc)
        except Exception as exc:
            log_error(f"Could not save {self.path.name}.", exc)
        try:
            tmp.unlink(missing_ok=True)
        except OSError:
            pass
        return False

    # ------------------------------------------------------------------- data
    def _row_for(self, d: date) -> int:
        row = self._rows.get(d)
        if row is None:
            row = self._ws.max_row + 1
            # skip any blank trailing row openpyxl may report
            while row > 2 and all(self._ws.cell(row - 1, c).value is None for c in range(1, 6)):
                row -= 1
            self._ws.cell(row, 1, d).number_format = "yyyy-mm-dd"
            for c in range(2, 6):
                self._ws.cell(row, c, 0)
            self._rows[d] = row
            self._dirty = True
        return row

    def add(self, stream: int, changes: dict[date, tuple[int, int]]) -> bool:
        """Add drained tracker changes for stream 1 or 2. Returns True if an
        event count changed (so the caller can save right away)."""
        if not self.is_open or not changes:
            return False
        ev_col, sec_col = _STREAM_COLS[stream]
        new_event = False
        for d in sorted(changes):
            events, seconds = changes[d]
            row = self._row_for(d)
            if events:
                cell = self._ws.cell(row, ev_col)
                cell.value = int(_as_number(cell.value)) + int(events)
                new_event = True
            if seconds:
                cell = self._ws.cell(row, sec_col)
                cell.value = int(_as_number(cell.value)) + int(seconds)
            self._dirty = True
        return new_event

    def ensure_row(self, d: date) -> None:
        if self.is_open:
            self._row_for(d)

    def blocked_seconds_between(self, first: date, last: date) -> tuple[int, int]:
        """(stream 1, stream 2) blocked seconds for dates first..last inclusive."""
        if not self.is_open:
            return 0, 0
        s1 = s2 = 0
        for d, r in self._rows.items():
            if first <= d <= last:
                s1 += int(_as_number(self._ws.cell(r, 3).value))
                s2 += int(_as_number(self._ws.cell(r, 5).value))
        return s1, s2

    def totals(self, d: date) -> tuple[int, int, int, int]:
        """(s1 events, s1 seconds, s2 events, s2 seconds) recorded for date d."""
        if not self.is_open or d not in self._rows:
            return 0, 0, 0, 0
        r = self._rows[d]
        return tuple(int(_as_number(self._ws.cell(r, c).value)) for c in range(2, 6))  # type: ignore


def read_blocked_seconds_between(path: Path, first: date, last: date) -> tuple[int, int]:
    """Same as ExcelLogger.blocked_seconds_between, but straight from the file
    (used while logging is stopped). Returns (0, 0) if the file can't be read."""
    path = Path(path)
    if not path.exists():
        return 0, 0
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return 0, 0
    try:
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]
        s1 = s2 = 0
        for row in ws.iter_rows(min_row=2, max_col=5, values_only=True):
            d = _as_date(row[0]) if row else None
            if d is not None and first <= d <= last:
                s1 += int(_as_number(row[2]))
                s2 += int(_as_number(row[4]))
        return s1, s2
    finally:
        wb.close()
